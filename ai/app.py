"""Private, stateless document intelligence. No hiring judgments or demographic extraction."""
import hashlib
import io
import json
import os
import re
import zipfile
from typing import Literal

import httpx
from docx import Document
from fastapi import Depends, FastAPI, HTTPException, Request
from pydantic import BaseModel, ConfigDict, Field
from pypdf import PdfReader

app = FastAPI(title="Workforce document intelligence", docs_url=None, redoc_url=None)
OLLAMA = os.getenv("OLLAMA_URL", "http://ollama:11434")
MODEL = os.getenv("OLLAMA_MODEL", "deepseek-r1:1.5b")
EMBED_MODEL = os.getenv("EMBED_MODEL", "nomic-embed-text")
PROMPT_VERSION = "evidence-gap-v5-scoped-references"
THINK = os.getenv("OLLAMA_THINK", "false").lower() == "true"
THREADS = max(1, min(32, int(os.getenv("OLLAMA_THREADS", "4"))))

async def private(request: Request):
    key = os.getenv("INTERNAL_API_KEY", "")
    if not key or not __import__('hmac').compare_digest(request.headers.get("x-internal-key", ""), key):
        raise HTTPException(401, "Private service")

class Strict(BaseModel):
    model_config = ConfigDict(extra="forbid")

class Section(Strict):
    id: str
    text: str

class Taxon(Strict):
    id: str
    name: str
    aliases: list[str]

class ExtractInput(Strict):
    sections: list[Section]
    skills: list[Taxon]

class Fact(Strict):
    kind: Literal["skill", "education", "experience", "technology"]
    value: str
    section_id: str
    quote: str
    skill_id: str | None = None
    mapping: str = "Unmapped"
    verification: str = "Unverified"

class Evidence(Strict):
    id: str
    skill_id: str | None
    value: str
    quote: str
    section_id: str
    verification: Literal["Documented", "Unverified", "Ambiguous", "Rejected"]
    proficiency: int | None = Field(default=None, ge=1, le=4)

class Need(Strict):
    skill_id: str
    name: str
    proficiency: int = Field(ge=1, le=4)

class AnalysisInput(Strict):
    requirements: list[Need]
    evidence: list[Evidence]

class Gap(Strict):
    skill_id: str
    status: Literal["Documented", "Missing", "Ambiguous", "Unverified"]
    evidence_ids: list[str]
    section_ids: list[str]
    explanation: str = Field(min_length=1, max_length=2000)
    uncertainty: str = Field(min_length=1, max_length=1000)
    proficiency_met: bool | None

class Output(Strict):
    gaps: list[Gap]

async def ollama(path: str, payload: dict, timeout: int = 180):
    try:
        async with httpx.AsyncClient(timeout=timeout) as client:
            r = await client.post(f"{OLLAMA}{path}", json=payload)
            r.raise_for_status()
            return r.json()
    except (httpx.HTTPError, ValueError) as exc:
        raise HTTPException(503, f"Local model unavailable or invalid response ({type(exc).__name__})") from exc

@app.get("/health")
async def health():
    return {"status": "ok", "model": MODEL, "embedding_model": EMBED_MODEL}

@app.get("/models", dependencies=[Depends(private)])
async def models():
    try:
        async with httpx.AsyncClient(timeout=5) as client:
            r = await client.get(f"{OLLAMA}/api/tags")
            r.raise_for_status()
            data = r.json()
        return {"available": True, "models": data.get("models", []), "configured_model": MODEL, "embedding_model": EMBED_MODEL}
    except (httpx.HTTPError, ValueError):
        return {"available": False, "models": [], "configured_model": MODEL, "embedding_model": EMBED_MODEL}

@app.post("/parse", dependencies=[Depends(private)])
async def parse(request: Request):
    # Bound stream before allocating or parsing attacker-controlled archives.
    chunks = bytearray()
    async for chunk in request.stream():
        chunks.extend(chunk)
        if len(chunks) > 5 * 1024 * 1024:
            raise HTTPException(413, "Document exceeds 5 MiB")
    raw = bytes(chunks)
    sections = []
    try:
        if raw.startswith(b"%PDF-"):
            reader = PdfReader(io.BytesIO(raw), strict=True)
            if reader.is_encrypted or len(reader.pages) > 40:
                raise ValueError("Encrypted or more than 40 pages")
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    sections.append(Section(id=f"page-{i + 1}", text=text))
        elif raw.startswith(b"PK"):
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(f.file_size for f in z.infolist()) > 20 * 1024 * 1024 or len(z.infolist()) > 200:
                    raise ValueError("Oversized archive")
                if "word/document.xml" not in z.namelist() or any("vbaProject" in x for x in z.namelist()):
                    raise ValueError("Only non-macro DOCX documents are supported")
            doc = Document(io.BytesIO(raw))
            texts = [p.text for p in doc.paragraphs if p.text.strip()]
            texts += [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
            sections = [Section(id=f"paragraph-{i + 1}", text=t) for i, t in enumerate(texts)]
        else:
            raise ValueError("Unsupported file signature")
        if not sections or sum(len(s.text) for s in sections) > 120000:
            raise ValueError("No selectable text or text exceeds 120,000 characters; OCR is not supported")
        return {"sections": [s.model_dump() for s in sections], "parser_version": "pypdf-5.5.0/docx-1.1.2"}
    except Exception as exc:
        raise HTTPException(422, "Unreadable, encrypted, scanned, or unsafe document") from exc

def extract_records(data: ExtractInput):
    facts = []
    for section in data.sections:
        for line in section.text.splitlines():
            line = line.strip()[:2000]
            if not line:
                continue
            for skill in data.skills:
                aliases = [skill.name, *skill.aliases]
                if any(re.search(r"(?<!\w)" + re.escape(a) + r"(?!\w)", line, re.I) for a in aliases):
                    facts.append(Fact(kind="skill", value=skill.name, section_id=section.id, quote=line,
                                      skill_id=skill.id, mapping="Alias", verification="Unverified"))
            # Explicit headings preserve only job-relevant facts; no demographic inference.
            for kind, pattern in [("education", r"^(education|qualification|degree)\s*:\s*(.+)"),
                                  ("experience", r"^(experience|employment|project)\s*:\s*(.+)"),
                                  ("technology", r"^(technology|technologies|skills)\s*:\s*(.+)")]:
                match = re.match(pattern, line, re.I)
                if match:
                    facts.append(Fact(kind=kind, value=match.group(2), section_id=section.id, quote=line))
    return [f.model_dump() for f in facts]

@app.post("/extract", dependencies=[Depends(private)])
async def extract(data: ExtractInput):
    return {"facts": extract_records(data), "extractor_version": "conservative-headings-alias-v1"}

@app.post("/embed", dependencies=[Depends(private)])
async def embed(data: dict):
    texts = data.get("texts")
    if not isinstance(texts, list) or not 1 <= len(texts) <= 100 or any(not isinstance(t, str) or len(t) > 4000 for t in texts):
        raise HTTPException(422, "1–100 texts of up to 4000 characters required")
    result = await ollama("/api/embed", {"model": EMBED_MODEL, "input": texts})
    vectors = result.get("embeddings", [])
    if len(vectors) != len(texts) or any(len(v) != 768 for v in vectors):
        raise HTTPException(502, "Embedding model must return 768-dimensional vectors")
    return {"vectors": vectors, "model": EMBED_MODEL}

# The model selects only IDs. Explanatory prose is constructed from checked evidence;
# unsupported model prose, extra attributes and hiring judgments are rejected by schema.
class Selection(Strict):
    skill_id: str
    evidence_ids: list[str]

class ModelOutput(Strict):
    selections: list[Selection]

def validate_and_ground(data: AnalysisInput, selected: ModelOutput) -> Output:
    needs = {n.skill_id: n for n in data.requirements}
    evidence = {e.id: e for e in data.evidence if e.verification != "Rejected"}
    if len(selected.selections) != len(needs) or {s.skill_id for s in selected.selections} != set(needs):
        raise ValueError("Requirements omitted or duplicated")
    gaps = []
    for selection in selected.selections:
        expected = {e.id for e in evidence.values() if e.skill_id == selection.skill_id}
        if set(selection.evidence_ids) != expected or len(selection.evidence_ids) != len(expected):
            raise ValueError("Unsupported or omitted evidence reference")
        facts = [evidence[eid] for eid in selection.evidence_ids]
        verified = [e for e in facts if e.verification == "Documented"]
        status = "Documented" if verified else "Ambiguous" if any(e.verification == "Ambiguous" for e in facts) else "Unverified" if facts else "Missing"
        levels = [e.proficiency for e in verified if e.proficiency is not None]
        met = max(levels) >= needs[selection.skill_id].proficiency if levels else None
        explanation = (f"{needs[selection.skill_id].name}: {len(verified)} human-documented record(s), {len(facts) - len(verified)} unverified or ambiguous record(s). Requirement level {needs[selection.skill_id].proficiency}."
                       if facts else f"No evidence for {needs[selection.skill_id].name} in the supplied records. Absence of evidence is not absence of skill.")
        uncertainty = "Proficiency is not established by reviewed evidence." if met is None else "Proficiency reflects a human-entered assessment, not an AI score."
        gaps.append(Gap(skill_id=selection.skill_id, status=status, evidence_ids=selection.evidence_ids,
                        section_ids=sorted({e.section_id for e in facts}), explanation=explanation, uncertainty=uncertainty, proficiency_met=met))
    return Output(gaps=gaps)

@app.post("/analyze", dependencies=[Depends(private)])
async def analyze(data: AnalysisInput):
    # Only canonical requirement labels and factual IDs/statuses enter inference. Resume text
    # and identity are withheld to minimize PII and prevent document prompt injection.
    # Compact opaque references reduce small-model copying errors and CPU latency.
    skill_refs = {n.skill_id: f"s{i + 1}" for i, n in enumerate(data.requirements)}
    relevant = sorted((e for e in data.evidence if e.verification != "Rejected" and e.skill_id in skill_refs),
                      key=lambda e: (skill_refs[e.skill_id], e.id))
    fact_refs = {e.id: f"e{i + 1}" for i, e in enumerate(relevant)}
    reverse_skills = {v: k for k, v in skill_refs.items()}
    reverse_facts = {v: k for k, v in fact_refs.items()}
    safe = {"requirements": [{"skill_id": skill_refs[n.skill_id], "name": n.name} for n in data.requirements],
            "evidence": [{"id": fact_refs[e.id], "skill_id": skill_refs[e.skill_id]}
                         for e in relevant]}
    prompt = "Return JSON only: keys are requirement skill_ids; values are lists of ALL matching evidence IDs. Match exact skill_id. Use [] if no evidence matches. No scores, recommendations or personal inference.\n" + json.dumps(safe, sort_keys=True)
    # A fixed set of JSON keys prevents a small model from looping over an unbounded
    # array of requirements. All semantic membership checks still run after inference.
    properties = {}
    for skill_id, compact_id in skill_refs.items():
        matching = [fact_refs[e.id] for e in data.evidence if e.id in fact_refs and e.skill_id == skill_id]
        properties[compact_id] = {"type": "array", "items": {"type": "string", "enum": matching or ["no-evidence"]},
                                  "minItems": len(matching), "maxItems": len(matching), "uniqueItems": True}
    schema = {"type": "object", "properties": properties, "required": list(properties), "additionalProperties": False}
    result = await ollama("/api/generate", {"model": MODEL, "prompt": prompt, "format": schema, "stream": False, "think": THINK,
                                           "options": {"temperature": 0, "num_predict": 512, "num_ctx": 2048, "num_thread": THREADS}})
    try:
        compact = json.loads(result["response"])
        if set(compact) != set(reverse_skills):
            raise ValueError("Invalid requirement keys")
        chosen = ModelOutput(selections=[Selection(skill_id=reverse_skills[sid],
                  evidence_ids=[reverse_facts[eid] for eid in eids]) for sid, eids in compact.items()])
        output = validate_and_ground(data, chosen)
    except (ValueError, KeyError) as exc:
        raise HTTPException(422, "Model response rejected: unsupported, incomplete, or invalid evidence references") from exc
    digest = None
    try:
        async with httpx.AsyncClient(timeout=5) as client:
            tags = (await client.get(f"{OLLAMA}/api/tags")).json().get("models", [])
        digest = next((m.get("digest") for m in tags if m.get("name") == MODEL), None)
    except (httpx.HTTPError, ValueError):
        pass
    return {"result": output.model_dump(), "metadata": {"model": result.get("model", MODEL), "model_digest": digest,
            "prompt_version": PROMPT_VERSION, "prompt_sha256": hashlib.sha256(prompt.encode()).hexdigest(),
            "thinking_enabled": THINK, "cpu_threads": THREADS, "context_tokens": 2048, "temperature": 0, "schema_version": "gap-v1", "grounding_version": "strict-ids-v1", "response_sha256": hashlib.sha256(result["response"].encode()).hexdigest(),
            "total_duration_ns": result.get("total_duration"), "generated_at": result.get("created_at")}}
